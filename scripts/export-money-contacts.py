#!/usr/bin/env python3
"""
从 money.xlsx 抽取内推联系方式，输出 CSV 供导入 Supabase。

money.xlsx 结构（第1行标题，第2行表头，第3行起数据）：
  A公司 B行业 C职位类别 D职位名称 E地区 F申请链接 G内推 H日期
  I工作职责&任职要求 J薪资 K内推方式 L来源

输出：supabase/money-contacts.csv
  列：source_url, apply_url, contact
"""

import csv
import os
import sys

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, 'money.xlsx')
OUT = os.path.join(ROOT, 'supabase', 'money-contacts.csv')

PLACEHOLDER = {'', '-', '无', 'nan', 'None'}


def main():
    if not os.path.exists(XLSX):
        print(f'❌ 找不到 {XLSX}')
        sys.exit(1)

    wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]

    rows = []
    total = 0
    for idx, r in enumerate(ws.iter_rows(min_row=3, values_only=True), start=3):
        if not r or len(r) < 12:
            continue
        total += 1
        apply_url = (r[5] or '').strip() if isinstance(r[5], str) else ''
        contact = (r[10] or '').strip() if isinstance(r[10], str) else ''
        source = (r[11] or '').strip() if isinstance(r[11], str) else ''
        if contact in PLACEHOLDER:
            continue
        rows.append((source, apply_url, contact))

    wb.close()

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(['source_url', 'apply_url', 'contact'])
        w.writerows(rows)

    print(f'✅ 扫描 {total} 行，抽出 {len(rows)} 条联系方式')
    print(f'   输出: supabase/money-contacts.csv')
    if rows:
        print('   样例:')
        for r in rows[:3]:
            print(f'     {r[0][:50]} | {r[2][:30]}')


if __name__ == '__main__':
    main()
